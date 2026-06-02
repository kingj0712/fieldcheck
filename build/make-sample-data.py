"""Generates sample FieldCheck import data (CSV + XLSX) with realistic BAS/controls
commissioning content. The CSV is what FieldCheck imports; the XLSX mirrors it for editing
in Excel (Save As -> CSV to import). Run:  python build/make-sample-data.py"""

import csv
import os

HEADERS = ["project_name", "checklist_name", "section", "item_text", "notes", "tags"]

P1 = "Riverside Medical Center - BAS Commissioning"
P2 = "Riverside Medical Center - Integration & Network"

# (project, checklist, section, item_text, notes, tags)
ROWS = [
    # ---- Project 1 / AHU-1 Checkout ----
    (P1, "AHU-1 Checkout", "Fan", "Verify supply fan command", "Command SF from BAS at 0, 50, and 100%; confirm VFD follows.", "AHU-1; BAS; Fan"),
    (P1, "AHU-1 Checkout", "Fan", "Verify supply fan status", "Confirm proof/status follows command; no false proof at 0%.", "AHU-1; Status"),
    (P1, "AHU-1 Checkout", "Fan", "Verify return fan tracking", "RF tracks SF per sequence; verify the tracking offset.", "AHU-1; Fan"),
    (P1, "AHU-1 Checkout", "Fan", "Verify VFD fault feedback", "Trip the VFD; confirm BAS alarm and fan-fail status.", "AHU-1; VFD; Safety"),
    (P1, "AHU-1 Checkout", "Sensors", "Verify SAT sensor", "Compare BAS supply-air temp to a calibrated field meter (+/- 1F).", "AHU-1; Sensor; Calibration"),
    (P1, "AHU-1 Checkout", "Sensors", "Verify MAT sensor", "Compare mixed-air temp to field reading.", "AHU-1; Sensor"),
    (P1, "AHU-1 Checkout", "Sensors", "Verify return air temp and humidity", "Check RAT and RH against a field instrument.", "AHU-1; Sensor"),
    (P1, "AHU-1 Checkout", "Sensors", "Verify duct static pressure", "Confirm DSP transmitter zero and span; compare to manometer.", "AHU-1; Sensor; Calibration"),
    (P1, "AHU-1 Checkout", "Dampers", "Verify OA damper stroke", "Command 0-100%; confirm full stroke and end switches.", "AHU-1; Damper"),
    (P1, "AHU-1 Checkout", "Dampers", "Verify economizer minimum position", "Confirm minimum OA position during occupied mode.", "AHU-1; Damper; Economizer"),
    (P1, "AHU-1 Checkout", "Dampers", "Verify RA/EA damper interlock", "OA opens as RA closes; check linkage and travel.", "AHU-1; Damper"),
    (P1, "AHU-1 Checkout", "Valves", "Verify cooling coil valve stroke", "Command CHW valve 0-100%; confirm actuator stroke.", "AHU-1; Valve"),
    (P1, "AHU-1 Checkout", "Valves", "Verify heating coil valve stroke", "Command HHW valve 0-100%; confirm actuator stroke.", "AHU-1; Valve"),
    (P1, "AHU-1 Checkout", "Safeties", "Verify low-limit freezestat", "Test freezestat; confirm it stops the fan and opens the HHW valve.", "AHU-1; Safety; High Priority"),
    (P1, "AHU-1 Checkout", "Safeties", "Verify smoke detector shutdown", "Activate duct smoke detector; confirm fan shutdown and alarm.", "AHU-1; Safety; High Priority"),
    (P1, "AHU-1 Checkout", "Safeties", "Verify high static cutout", "Confirm the high duct-static safety trips the fan.", "AHU-1; Safety"),

    # ---- Project 1 / VAV Boxes - Floor 3 ----
    (P1, "VAV Boxes - Floor 3", "Airflow", "Verify VAV minimum airflow", "Confirm minimum CFM at zero load matches the schedule.", "VAV; Airflow; Floor 3"),
    (P1, "VAV Boxes - Floor 3", "Airflow", "Verify VAV maximum airflow", "Confirm maximum CFM on a call for cooling.", "VAV; Airflow; Floor 3"),
    (P1, "VAV Boxes - Floor 3", "Airflow", "Verify airflow sensor calibration", "Check K-factor / velocity pressure against the balancer reading.", "VAV; Calibration"),
    (P1, "VAV Boxes - Floor 3", "Reheat", "Verify reheat valve stroke", "Command HHW reheat 0-100%; confirm discharge-air temp rise.", "VAV; Valve; Reheat"),
    (P1, "VAV Boxes - Floor 3", "Reheat", "Verify discharge air temp sensor", "Compare DAT to a field reading.", "VAV; Sensor"),
    (P1, "VAV Boxes - Floor 3", "Control", "Verify zone temp sensor", "Compare space temp to a field thermometer.", "VAV; Sensor"),
    (P1, "VAV Boxes - Floor 3", "Control", "Verify occupancy override", "Press the override; confirm timed occupied mode.", "VAV; Occupancy"),

    # ---- Project 1 / Graphics Review ----
    (P1, "Graphics Review", "Navigation", "Verify AHU graphic links", "Confirm the graphic opens the correct equipment view.", "Graphics; AHU-1"),
    (P1, "Graphics Review", "Navigation", "Verify floor plan navigation", "Confirm the floor plan links to each VAV box.", "Graphics; Navigation"),
    (P1, "Graphics Review", "Trends", "Verify SAT trend logging", "Confirm SAT history is collecting at a 5-minute interval.", "Graphics; Trend"),
    (P1, "Graphics Review", "Trends", "Verify runtime totalization", "Confirm fan runtime is accumulating.", "Graphics; Trend"),
    (P1, "Graphics Review", "Alarms", "Verify alarm console link", "Confirm the alarm console opens and filters by AHU.", "Graphics; Alarm"),

    # ---- Project 2 / BACnet Integration ----
    (P2, "BACnet Integration", "Network", "Verify BBMD configuration", "Confirm BBMD/BDT entries are correct across subnets.", "BACnet; Network; Niagara"),
    (P2, "BACnet Integration", "Network", "Verify device discovery", "All controllers discovered; no duplicate device instance IDs.", "BACnet; Network"),
    (P2, "BACnet Integration", "Network", "Verify COV subscriptions", "Confirm COV is active for key points; no polling fallback.", "BACnet; Network"),
    (P2, "BACnet Integration", "Points", "Verify point mapping", "Spot-check 10% of points map to the correct objects.", "BACnet; Points"),
    (P2, "BACnet Integration", "Points", "Verify units and scaling", "Confirm engineering units and scaling are correct.", "BACnet; Points; Calibration"),

    # ---- Project 2 / Alarm Testing ----
    (P2, "Alarm Testing", "Critical Alarms", "Verify space temp high/low alarm", "Force an out-of-range value; confirm alarm and priority.", "Alarm; High Priority"),
    (P2, "Alarm Testing", "Critical Alarms", "Verify equipment fail alarm", "Simulate a fan failure; confirm a critical alarm.", "Alarm; High Priority"),
    (P2, "Alarm Testing", "Notifications", "Verify email notification routing", "Confirm alarm emails reach the on-call group.", "Alarm; Notification"),
    (P2, "Alarm Testing", "Notifications", "Verify alarm acknowledgment", "Acknowledge an alarm; confirm state change and audit entry.", "Alarm"),
]

out_dir = os.path.join(os.path.dirname(__file__), "..", "samples")
out_dir = os.path.abspath(out_dir)
os.makedirs(out_dir, exist_ok=True)
csv_path = os.path.join(out_dir, "commissioning-sample.csv")
xlsx_path = os.path.join(out_dir, "commissioning-sample.xlsx")

# ---- CSV (this is what FieldCheck imports) ----
with open(csv_path, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
    w.writerow(HEADERS)
    w.writerows(ROWS)

# ---- XLSX (for editing in Excel; Save As -> CSV to import) ----
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "FieldCheck Import"

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="2E6E8E")
ws.append(HEADERS)
for col, _ in enumerate(HEADERS, start=1):
    c = ws.cell(row=1, column=col)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(vertical="center")

for row in ROWS:
    ws.append(list(row))

widths = [34, 22, 16, 40, 52, 26]
for i, width in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = width

ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 20
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=4).alignment = Alignment(wrap_text=False, vertical="center")
    ws.cell(row=r, column=5).alignment = Alignment(wrap_text=True, vertical="center")
ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

wb.save(xlsx_path)

projects = sorted({r[0] for r in ROWS})
checklists = sorted({(r[0], r[1]) for r in ROWS})
print(f"Wrote {csv_path}")
print(f"Wrote {xlsx_path}")
print(f"Rows: {len(ROWS)} items | Projects: {len(projects)} | Checklists: {len(checklists)}")
